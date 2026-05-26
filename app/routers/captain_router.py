"""Captain / On board — SOF events, ETA shifts, onboard messages, cargo docs.

Reprises de la V3.0.0 :
- Saisie chronologique d'événements SOF (EOSP, SOSP, NOR, PILOT_ON…).
- Déclaration d'un décalage d'ETA avec motif obligatoire (9 raisons codifiées).
- Messagerie de bord avec @mentions et bot MYTOWT_BOT.
- Génération de cargo documents (NOR, NOR_RT, LOP, Mate's Receipt).
"""
from __future__ import annotations

import re
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.database import get_db
from app.models.finance import PortConfig
from app.models.leg import Leg
from app.models.noon_report import NoonReport
from app.models.port import Port
from app.models.sof_event import (
    CargoDocument, ETA_SHIFT_REASONS, EtaShift,
    OnboardMessage, OnboardMessageMention, SOF_EVENT_TYPES, SofEvent,
)
from app.models.user import User
from app.models.vessel import Vessel
from app.models.watch_log import WatchLog
from app.permissions import require_permission
from app.services import weather as wx
from app.services.activity import record as activity_record
from app.services.signature import (
    compute_noon_hash, compute_sof_hash, compute_watch_hash, sign_record,
)
from app.templating import templates

router = APIRouter(prefix="/captain", tags=["captain"])

MENTION_RE = re.compile(r"@([A-Za-z0-9_]{2,40})")
BOT_TRIGGERS = ("@MYTOWT_BOT", "@mytowt_bot", "@bot")


@router.get("", response_class=HTMLResponse)
@router.get("/", response_class=HTMLResponse)
async def captain_index(
    request: Request,
    leg_id: int | None = None,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "C")),
) -> HTMLResponse:
    legs = list((await db.execute(
        select(Leg).order_by(Leg.etd.desc()).limit(20)
    )).scalars().all())
    selected = (await db.get(Leg, leg_id)) if leg_id else (legs[0] if legs else None)
    events: list[SofEvent] = []
    eta_shifts: list[EtaShift] = []
    messages: list[OnboardMessage] = []
    docs: list[CargoDocument] = []
    vessel = None
    if selected:
        events = list((await db.execute(
            select(SofEvent).where(SofEvent.leg_id == selected.id)
            .order_by(SofEvent.occurred_at.desc()).limit(100)
        )).scalars().all())
        eta_shifts = list((await db.execute(
            select(EtaShift).where(EtaShift.leg_id == selected.id)
            .order_by(EtaShift.declared_at.desc())
        )).scalars().all())
        messages = list((await db.execute(
            select(OnboardMessage).where(OnboardMessage.leg_id == selected.id)
            .order_by(OnboardMessage.created_at.desc()).limit(50)
        )).scalars().all())
        docs = list((await db.execute(
            select(CargoDocument).where(CargoDocument.leg_id == selected.id)
            .order_by(CargoDocument.issued_at.desc())
        )).scalars().all())
        vessel = await db.get(Vessel, selected.vessel_id)
    return templates.TemplateResponse(
        "staff/captain/index.html",
        {
            "request": request, "user": user,
            "legs": legs, "leg": selected, "vessel": vessel,
            "events": events, "eta_shifts": eta_shifts,
            "messages": messages, "docs": docs,
            "event_types": SOF_EVENT_TYPES,
            "eta_reasons": ETA_SHIFT_REASONS,
        },
    )


@router.post("/legs/{leg_id}/sof")
async def add_sof_event(
    leg_id: int,
    request: Request,
    event_type: str = Form(...),
    occurred_at: str = Form(...),
    label: str | None = Form(None),
    latitude: float | None = Form(None),
    longitude: float | None = Form(None),
    notes: str | None = Form(None),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    if event_type not in SOF_EVENT_TYPES:
        raise HTTPException(status_code=400, detail="invalid event_type")
    if not await db.get(Leg, leg_id):
        raise HTTPException(status_code=404)
    e = SofEvent(
        leg_id=leg_id,
        event_type=event_type,
        label=label,
        occurred_at=datetime.fromisoformat(occurred_at),
        latitude=latitude, longitude=longitude,
        notes=notes,
        recorded_by_id=user.id,
        recorded_by_name=user.full_name or user.username,
    )
    db.add(e)
    await db.flush()
    await activity_record(
        db, action="create", user_id=user.id, user_name=user.full_name or user.username,
        user_role=user.role, module="captain", entity_type="sof_event",
        entity_id=e.id, entity_label=f"{event_type}@{occurred_at}",
        ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


@router.post("/legs/{leg_id}/eta-shift")
async def declare_eta_shift(
    leg_id: int,
    request: Request,
    new_eta: str = Form(...),
    reason: str = Form(...),
    detail: str | None = Form(None),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    if reason not in ETA_SHIFT_REASONS:
        raise HTTPException(status_code=400, detail="invalid reason")
    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)
    shift = EtaShift(
        leg_id=leg_id,
        previous_eta=leg.eta,
        new_eta=datetime.fromisoformat(new_eta),
        reason=reason, detail=detail,
        declared_by_id=user.id,
        declared_by_name=user.full_name or user.username,
    )
    db.add(shift)
    leg.eta = shift.new_eta
    await db.flush()
    await activity_record(
        db, action="update", user_id=user.id, user_name=user.full_name or user.username,
        user_role=user.role, module="captain", entity_type="eta_shift",
        entity_id=shift.id, entity_label=f"leg={leg_id} reason={reason}",
        ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


@router.post("/legs/{leg_id}/messages")
async def post_onboard_message(
    leg_id: int,
    request: Request,
    body: str = Form(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)
    msg = OnboardMessage(
        leg_id=leg_id, vessel_id=leg.vessel_id,
        author_id=user.id,
        author_name=user.full_name or user.username,
        is_bot=False, body=body.strip(),
    )
    db.add(msg)
    await db.flush()
    # Detect @mentions
    for tag in MENTION_RE.findall(body):
        target_user = (await db.execute(
            select(User).where(User.username == tag)
        )).scalar_one_or_none()
        db.add(OnboardMessageMention(
            message_id=msg.id,
            mentioned_user_id=target_user.id if target_user else None,
            mentioned_text=tag,
        ))
    # Bot reply (placeholder — extended by chat service in Phase 5)
    if any(t.lower() in body.lower() for t in BOT_TRIGGERS):
        db.add(OnboardMessage(
            leg_id=leg_id, vessel_id=leg.vessel_id,
            author_id=None, author_name="MYTOWT_BOT",
            is_bot=True,
            body=f"Bonjour {user.full_name or user.username}, le bot Kairos est en cours d'intégration.",
        ))
    await db.flush()
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


@router.post("/legs/{leg_id}/docs")
async def create_cargo_document(
    leg_id: int,
    request: Request,
    kind: str = Form(...),
    reference: str | None = Form(None),
    issued_at: str = Form(...),
    party_name: str | None = Form(None),
    body: str | None = Form(None),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    if not await db.get(Leg, leg_id):
        raise HTTPException(status_code=404)
    d = CargoDocument(
        leg_id=leg_id, kind=kind, reference=reference,
        issued_at=datetime.fromisoformat(issued_at),
        party_name=party_name, body=body,
    )
    db.add(d)
    await db.flush()
    await activity_record(
        db, action="create", user_id=user.id, user_name=user.full_name or user.username,
        user_role=user.role, module="captain", entity_type="cargo_document",
        entity_id=d.id, entity_label=f"{kind} {reference or ''}".strip(),
        ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


# ─────────────────────────────────────────────────────────────────────
#                  Prochaine escale — vue commandant
# ─────────────────────────────────────────────────────────────────────


@router.get("/next-port", response_class=HTMLResponse)
async def next_port(
    request: Request,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "C")),
) -> HTMLResponse:
    """Synthèse "prochaine escale" — port d'arrivée du prochain leg actif.

    Affiche : nom port, ETA, distance restante, contacts (PortConfig),
    météo forecast au moment ETA, événements SOF récents. Filtré par
    ``user.assigned_vessel_id`` si renseigné.
    """
    from datetime import datetime, timezone
    now = datetime.now(timezone.utc)

    # Prochain leg actif (ATD posé, pas encore arrivé) ou prochain ETD.
    stmt_active = (
        select(Leg)
        .where(Leg.atd.is_not(None))
        .where(Leg.ata.is_(None))
        .order_by(Leg.etd.asc())
        .limit(1)
    )
    stmt_planned = (
        select(Leg)
        .where(Leg.etd > now)
        .where(Leg.ata.is_(None))
        .order_by(Leg.etd.asc())
        .limit(1)
    )
    if getattr(user, "assigned_vessel_id", None):
        stmt_active = stmt_active.where(Leg.vessel_id == user.assigned_vessel_id)
        stmt_planned = stmt_planned.where(Leg.vessel_id == user.assigned_vessel_id)

    leg = (await db.execute(stmt_active)).scalar_one_or_none()
    if leg is None:
        leg = (await db.execute(stmt_planned)).scalar_one_or_none()

    pod = None
    pod_config: PortConfig | None = None
    vessel = None
    weather_point = None
    sof_recent: list[SofEvent] = []
    if leg is not None:
        pod = await db.get(Port, leg.arrival_port_id)
        vessel = await db.get(Vessel, leg.vessel_id)
        if pod:
            pod_config = (await db.execute(
                select(PortConfig).where(PortConfig.port_id == pod.id)
            )).scalar_one_or_none()
        if pod and pod.latitude is not None and pod.longitude is not None and leg.eta:
            try:
                weather_point = await wx.fetch_at(pod.latitude, pod.longitude, leg.eta)
            except Exception:
                weather_point = None
        sof_recent = list((await db.execute(
            select(SofEvent).where(SofEvent.leg_id == leg.id)
            .order_by(SofEvent.occurred_at.desc()).limit(8)
        )).scalars().all())

    return templates.TemplateResponse(
        "staff/captain/next_port.html",
        {
            "request": request, "user": user,
            "leg": leg, "vessel": vessel, "pod": pod, "pod_config": pod_config,
            "weather_point": weather_point,
            "weather_summary": wx.summarize(weather_point),
            "sof_recent": sof_recent,
            "now": now,
        },
    )


# ─────────────────────────────────────────────────────────────────────
#                  Signature / lock — SOF / noon / watch
# ─────────────────────────────────────────────────────────────────────


@router.post("/sof-events/{event_id}/sign")
async def sign_sof_event(
    event_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    """Signe un SOF event → ``is_locked = True``, plus de modification possible."""
    e = await db.get(SofEvent, event_id)
    if e is None:
        raise HTTPException(status_code=404)
    sign_record(e, user, hash_fn=compute_sof_hash)
    await db.flush()
    await activity_record(
        db, action="sign", user_id=user.id, user_name=user.full_name or user.username,
        user_role=user.role, module="captain", entity_type="sof_event",
        entity_id=e.id, entity_label=f"{e.event_type}@{e.occurred_at.isoformat()}",
        detail=e.signature_hash[:12] if e.signature_hash else None,
        ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={e.leg_id}", status_code=303)


@router.post("/noon-reports/{report_id}/sign")
async def sign_noon_report(
    report_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    n = await db.get(NoonReport, report_id)
    if n is None:
        raise HTTPException(status_code=404)
    sign_record(n, user, hash_fn=compute_noon_hash)
    await db.flush()
    await activity_record(
        db, action="sign", user_id=user.id, user_name=user.full_name or user.username,
        user_role=user.role, module="captain", entity_type="noon_report",
        entity_id=n.id, entity_label=f"leg={n.leg_id}@{n.recorded_at.isoformat()}",
        detail=n.signature_hash[:12] if n.signature_hash else None,
        ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/onboard/navigation?leg_id={n.leg_id}", status_code=303)


@router.post("/watch-logs/{log_id}/sign")
async def sign_watch_log(
    log_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    w = await db.get(WatchLog, log_id)
    if w is None:
        raise HTTPException(status_code=404)
    sign_record(w, user, hash_fn=compute_watch_hash)
    await db.flush()
    await activity_record(
        db, action="sign", user_id=user.id, user_name=user.full_name or user.username,
        user_role=user.role, module="captain", entity_type="watch_log",
        entity_id=w.id, entity_label=f"leg={w.leg_id} {w.watch_date} {w.watch_period}",
        detail=w.signature_hash[:12] if w.signature_hash else None,
        ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/onboard/navigation?leg_id={w.leg_id}", status_code=303)


@router.get("/legs/{leg_id}/sof.pdf")
async def captain_sof_pdf(
    leg_id: int,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "C")),
):
    """Génère le SOF commandant (SofEvent) en PDF WeasyPrint."""
    from datetime import timezone

    from fastapi.responses import Response
    from weasyprint import HTML  # local import — heavy native deps

    from app.config import settings

    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)

    vessel = await db.get(Vessel, leg.vessel_id) if leg.vessel_id else None
    pol = await db.get(Port, leg.departure_port_id) if leg.departure_port_id else None
    pod = await db.get(Port, leg.arrival_port_id) if leg.arrival_port_id else None

    events = list((await db.execute(
        select(SofEvent)
        .where(SofEvent.leg_id == leg_id)
        .order_by(SofEvent.occurred_at.asc())
    )).scalars().all())

    tpl = templates.get_template("pdf/sof_captain.html")
    html = tpl.render(
        leg=leg,
        vessel=vessel,
        pol=pol,
        pod=pod,
        events=events,
        issued_at=datetime.now(timezone.utc),
        site_url=settings.site_url,
    )
    pdf = HTML(string=html, base_url=settings.site_url).write_pdf()
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="SOF_{leg.leg_code}.pdf"'
        },
    )


@router.get("/legs/{leg_id}/sof.xlsx")
async def captain_sof_xlsx(
    leg_id: int,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "C")),
):
    """Exporte le SOF commandant (SofEvent + ETA shifts) en classeur Excel."""
    import io
    from datetime import timezone

    import openpyxl
    from fastapi.responses import Response

    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)

    events = list((await db.execute(
        select(SofEvent)
        .where(SofEvent.leg_id == leg_id)
        .order_by(SofEvent.occurred_at.asc())
    )).scalars().all())

    eta_shifts = list((await db.execute(
        select(EtaShift)
        .where(EtaShift.leg_id == leg_id)
        .order_by(EtaShift.declared_at.asc())
    )).scalars().all())

    wb = openpyxl.Workbook()

    # ── Sheet 1: SOF Events ──────────────────────────────────────────
    ws_sof = wb.active
    ws_sof.title = "SOF Events"
    ws_sof.append([
        "#", "Type", "Label",
        "Occurred At (UTC)", "Port", "Lat", "Lon",
        "Notes", "Signed", "Signed By", "Signed At",
    ])
    for ev in events:
        ws_sof.append([
            ev.id,
            ev.event_type,
            ev.label or "",
            ev.occurred_at.strftime("%Y-%m-%d %H:%M") if ev.occurred_at else "",
            "",  # port_id — no eager-loaded name available without extra query
            ev.latitude if ev.latitude is not None else "",
            ev.longitude if ev.longitude is not None else "",
            ev.notes or "",
            "Oui" if ev.is_locked else "Non",
            ev.signed_by_name or "",
            ev.signed_at.strftime("%Y-%m-%d %H:%M") if ev.signed_at else "",
        ])

    # ── Sheet 2: ETA Shifts ──────────────────────────────────────────
    ws_eta = wb.create_sheet(title="ETA Shifts")
    ws_eta.append([
        "Declared At", "Reason", "New ETA", "Delta Hours", "Notes",
    ])
    for shift in eta_shifts:
        delta_h = ""
        if shift.previous_eta and shift.new_eta:
            delta_td = shift.new_eta - shift.previous_eta
            delta_h = round(delta_td.total_seconds() / 3600, 2)
        ws_eta.append([
            shift.declared_at.strftime("%Y-%m-%d %H:%M") if shift.declared_at else "",
            shift.reason,
            shift.new_eta.strftime("%Y-%m-%d %H:%M") if shift.new_eta else "",
            delta_h,
            shift.detail or "",
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="SOF_{leg.leg_code}.xlsx"'
        },
    )


# ──────────────────────────────────────────────────────────────────────────────
# A — Cargo document PDF generation
# ──────────────────────────────────────────────────────────────────────────────

_DOC_TEMPLATES: dict[str, str] = {
    "NOR": "pdf/cargo_doc_nor.html",
    "NOR_RT": "pdf/cargo_doc_nor.html",
    "LOP_GENERAL": "pdf/cargo_doc_lop.html",
    "LOP_DRAFT": "pdf/cargo_doc_lop.html",
    "MATES_RECEIPT": "pdf/cargo_doc_mates_receipt.html",
    "OTHER": "pdf/cargo_doc_nor.html",  # fallback layout
}


@router.get("/legs/{leg_id}/docs/{doc_id}.pdf")
async def captain_cargo_doc_pdf(
    leg_id: int,
    doc_id: int,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "C")),
):
    from fastapi.responses import Response
    from weasyprint import HTML

    from app.config import settings

    doc = (await db.execute(
        select(CargoDocument).where(
            CargoDocument.id == doc_id, CargoDocument.leg_id == leg_id
        )
    )).scalar_one_or_none()
    if doc is None:
        raise HTTPException(status_code=404)

    leg = await db.get(Leg, leg_id)
    vessel = await db.get(Vessel, leg.vessel_id) if leg and leg.vessel_id else None
    pol = await db.get(Port, leg.departure_port_id) if leg and leg.departure_port_id else None
    pod = await db.get(Port, leg.arrival_port_id) if leg and leg.arrival_port_id else None

    tpl_name = _DOC_TEMPLATES.get(doc.kind, "pdf/cargo_doc_nor.html")
    tpl = templates.get_template(tpl_name)
    html = tpl.render(
        doc=doc, leg=leg, vessel=vessel, pol=pol, pod=pod,
        issued_at=doc.issued_at,
        site_url=settings.site_url,
    )
    pdf = HTML(string=html, base_url=settings.site_url).write_pdf()
    safe_ref = (doc.reference or str(doc.id)).replace("/", "-")
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{doc.kind}_{safe_ref}.pdf"'},
    )


# ──────────────────────────────────────────────────────────────────────────────
# D — Pièces jointes aux cargo documents
# ──────────────────────────────────────────────────────────────────────────────

@router.post("/legs/{leg_id}/docs/{doc_id}/attach")
async def attach_cargo_doc(
    leg_id: int,
    doc_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    from app.services.safe_files import UploadRejected, save_upload

    doc = (await db.execute(
        select(CargoDocument).where(
            CargoDocument.id == doc_id, CargoDocument.leg_id == leg_id
        )
    )).scalar_one_or_none()
    if doc is None:
        raise HTTPException(status_code=404)

    form = await request.form()
    upload = form.get("file")
    if upload is None or not hasattr(upload, "read"):
        raise HTTPException(status_code=422, detail="Fichier manquant")

    content = await upload.read()
    try:
        rel_path, _ = save_upload(content, upload.filename or "attachment", subdir="captain_docs")
    except UploadRejected as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    doc.file_path = rel_path
    await db.flush()
    await activity_record(
        db, action="cargo_doc_attach",
        user_id=user.id, user_name=user.username, user_role=user.role,
        module="captain", entity_type="cargo_document", entity_id=doc.id,
        detail=upload.filename, ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


@router.get("/legs/{leg_id}/docs/{doc_id}/attachment")
async def download_cargo_doc_attachment(
    leg_id: int,
    doc_id: int,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "C")),
):
    from fastapi.responses import FileResponse
    from app.services.safe_files import UploadRejected, resolve_path

    doc = (await db.execute(
        select(CargoDocument).where(
            CargoDocument.id == doc_id, CargoDocument.leg_id == leg_id
        )
    )).scalar_one_or_none()
    if doc is None or not doc.file_path:
        raise HTTPException(status_code=404)

    try:
        path = resolve_path(doc.file_path)
    except UploadRejected:
        raise HTTPException(status_code=400)

    return FileResponse(path=str(path), filename=path.name)


# ──────────────────────────────────────────────────────────────────────────────
# B — Workflow de clôture de voyage
# ──────────────────────────────────────────────────────────────────────────────

@router.post("/legs/{leg_id}/closure/submit")
async def closure_submit(
    leg_id: int,
    request: Request,
    notes: str | None = Form(None),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)
    if leg.closure_submitted_at:
        raise HTTPException(status_code=400, detail="Clôture déjà soumise")

    leg.closure_submitted_at = datetime.now(timezone.utc)
    leg.closure_submitted_by = user.username
    if notes:
        leg.closure_notes = notes
    await db.flush()

    try:
        from app.services.notifications import create as notif_create
        await notif_create(
            db, type="info",
            title=f"Clôture soumise — {leg.leg_code}",
            link=f"/captain?leg_id={leg_id}",
            target_role="operation",
        )
    except Exception:
        pass

    await activity_record(
        db, action="voyage_closure_submit",
        user_id=user.id, user_name=user.username, user_role=user.role,
        module="captain", entity_type="leg", entity_id=leg.id,
        entity_label=leg.leg_code, ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


@router.post("/legs/{leg_id}/closure/review")
async def closure_review(
    leg_id: int,
    request: Request,
    notes: str | None = Form(None),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "M")),
):
    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)
    if not leg.closure_submitted_at:
        raise HTTPException(status_code=400, detail="Clôture non encore soumise")
    if leg.closure_reviewed_at:
        raise HTTPException(status_code=400, detail="Déjà validée")

    leg.closure_reviewed_at = datetime.now(timezone.utc)
    leg.closure_reviewed_by = user.username
    if notes:
        leg.closure_notes = (leg.closure_notes or "") + f"\n[Validation opérations] {notes}"
    await db.flush()

    try:
        from app.services.notifications import create as notif_create
        await notif_create(
            db, type="info",
            title=f"Clôture validée opérations — {leg.leg_code}",
            link=f"/captain?leg_id={leg_id}",
            target_role="manager_maritime",
        )
    except Exception:
        pass

    await activity_record(
        db, action="voyage_closure_review",
        user_id=user.id, user_name=user.username, user_role=user.role,
        module="captain", entity_type="leg", entity_id=leg.id,
        entity_label=leg.leg_code, ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


@router.post("/legs/{leg_id}/closure/approve")
async def closure_approve(
    leg_id: int,
    request: Request,
    notes: str | None = Form(None),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_permission("captain", "S")),
):
    leg = await db.get(Leg, leg_id)
    if leg is None:
        raise HTTPException(status_code=404)
    if not leg.closure_reviewed_at:
        raise HTTPException(status_code=400, detail="Validation opérations requise d'abord")
    if leg.closure_approved_at:
        raise HTTPException(status_code=400, detail="Déjà approuvée")

    leg.closure_approved_at = datetime.now(timezone.utc)
    if notes:
        leg.closure_notes = (leg.closure_notes or "") + f"\n[Approbation] {notes}"
    leg.status = "completed"
    await db.flush()

    try:
        from app.services.kpi import compute_for_leg
        await compute_for_leg(db, leg)
    except Exception:
        pass

    await activity_record(
        db, action="voyage_closure_approve",
        user_id=user.id, user_name=user.username, user_role=user.role,
        module="captain", entity_type="leg", entity_id=leg.id,
        entity_label=leg.leg_code, ip_address=_client_ip(request),
    )
    return RedirectResponse(url=f"/captain?leg_id={leg_id}", status_code=303)


def _client_ip(request: Request) -> str | None:
    return request.headers.get("x-forwarded-for") or (request.client.host if request.client else None)
